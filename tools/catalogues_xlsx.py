# -*- coding: utf-8 -*-
"""Génère docs/ROADLINE_Vehicules.xlsx et docs/ROADLINE_Armes.xlsx (catalogues lisibles pour le fondateur).
Sources : tools/data/vehicules_concession.json (concession, aéroport, marina), server/locales-fr/weapons.json (noms
français), server/overrides/ox_inventory/data/shops.lua (Ammu-Nation), gs_blackmarket/shared/config.lua (marché noir)
et, pour les poids / calibres, data/weapons.lua d'ox_inventory (chemin en argument, optionnel).
Usage : python tools/catalogues_xlsx.py [chemin/vers/ox_inventory/data/weapons.lua]"""
import io, json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

FONT = 'Arial'
HEAD_FILL = PatternFill('solid', fgColor='2A1440')
BAND_FILL = PatternFill('solid', fgColor='F6F1FC')
THIN = Side(style='thin', color='D9CFE8')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0 "$"'

def sheet_table(ws, headers, rows, widths, money_cols=(), note=None):
    """Tableau : en-tête violet, bandes, filtre, volets figés, largeurs."""
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border = BORDER
    for r in rows:
        ws.append(list(r))
    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)), start=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            if i % 2 == 0: cell.fill = BAND_FILL
            if cell.column in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = MONEY
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{ws.max_row}'
    ws.row_dimensions[1].height = 30
    if note:
        ws.cell(row=ws.max_row + 2, column=1, value=note).font = Font(name=FONT, italic=True, size=9, color='666666')

def summary_sheet(wb, title, intro, lines):
    ws = wb.create_sheet(title)
    ws['A1'] = intro
    ws['A1'].font = Font(name=FONT, bold=True, size=12)
    r = 3
    for label, formula, fmt in lines:
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, size=10)
        c = ws.cell(row=r, column=2, value=formula)
        c.font = Font(name=FONT, size=10, bold=True)
        if fmt: c.number_format = fmt
        r += 1
    ws.column_dimensions['A'].width = 46
    ws.column_dimensions['B'].width = 18
    ws.cell(row=r + 1, column=1, value='Les chiffres de cette feuille sont des formules (COUNTIF / SUMIF…) calculées sur la feuille de données : '
            'ils se mettent à jour si tu modifies le tableau.').font = Font(name=FONT, italic=True, size=9, color='666666')
    return ws

# --------------------------------------------------------------------------------------------- véhicules
def vehicules():
    d = json.load(io.open('tools/data/vehicules_concession.json', encoding='utf-8'))
    SHOP = {'pdm': 'Concession', 'boats': 'Marina', 'air': 'Aéroport'}
    rows = []
    for shop, cats in d['shops'].items():
        for cat, lst in cats.items():
            for v in lst:
                rows.append((SHOP.get(shop, shop), cat, v.get('brand', ''), v.get('name', ''), v.get('model', ''), v.get('price', 0)))
    rows.sort(key=lambda r: (r[0], r[1], r[5], r[3]))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Véhicules'
    sheet_table(ws, ['Boutique', 'Catégorie', 'Marque', 'Véhicule', 'Modèle (spawn)', 'Prix'], rows,
                [14, 14, 16, 26, 18, 13], money_cols=(6,),
                note='Source : ' + d.get('source', '') + ' · généré par tools/catalogues_xlsx.py')
    n = len(rows)
    lines = [('Véhicules en vente', f"=COUNTA('Véhicules'!E2:E{n + 1})", '0'),
             ('Concession', f"=COUNTIF('Véhicules'!A2:A{n + 1},\"Concession\")", '0'),
             ('Aéroport', f"=COUNTIF('Véhicules'!A2:A{n + 1},\"Aéroport\")", '0'),
             ('Marina', f"=COUNTIF('Véhicules'!A2:A{n + 1},\"Marina\")", '0'),
             ('Prix le plus bas', f"=MIN('Véhicules'!F2:F{n + 1})", MONEY),
             ('Prix le plus haut', f"=MAX('Véhicules'!F2:F{n + 1})", MONEY),
             ('Prix moyen', f"=AVERAGE('Véhicules'!F2:F{n + 1})", MONEY),
             ('Valeur totale du catalogue', f"=SUM('Véhicules'!F2:F{n + 1})", MONEY)]
    cats = sorted(set(r[1] for r in rows))
    lines.append(('', None, None))
    for c in cats:
        lines.append((f'Catégorie « {c} »', f"=COUNTIF('Véhicules'!B2:B{n + 1},\"{c}\")", '0'))
    summary_sheet(wb, 'Résumé', 'RoadLine RP · Catalogue des véhicules (résumé)', lines)
    wb.calculation.fullCalcOnLoad = True
    wb.save('docs/ROADLINE_Vehicules.xlsx')
    return n

# --------------------------------------------------------------------------------------------- armes
MELEE = ('KNIFE', 'NIGHTSTICK', 'HAMMER', 'BAT', 'CROWBAR', 'GOLFCLUB', 'BOTTLE', 'DAGGER', 'HATCHET', 'KNUCKLE', 'MACHETE', 'SWITCHBLADE',
         'BATTLEAXE', 'POOLCUE', 'WRENCH', 'STONE_HATCHET', 'FLASHLIGHT', 'CANDYCANE')
PISTOLS = ('PISTOL', 'COMBATPISTOL', 'APPISTOL', 'STUNGUN', 'PISTOL50', 'SNSPISTOL', 'HEAVYPISTOL', 'VINTAGEPISTOL', 'FLAREGUN', 'MARKSMANPISTOL',
           'REVOLVER', 'DOUBLEACTION', 'RAYPISTOL', 'CERAMICPISTOL', 'NAVYREVOLVER', 'GADGETPISTOL', 'PISTOLXM3', 'STUNGUN_MP')
SMG = ('MICROSMG', 'SMG', 'ASSAULTSMG', 'COMBATPDW', 'MACHINEPISTOL', 'MINISMG', 'RAYCARBINE', 'TECPISTOL')
SHOTGUNS = ('PUMPSHOTGUN', 'SAWNOFFSHOTGUN', 'ASSAULTSHOTGUN', 'BULLPUPSHOTGUN', 'MUSKET', 'HEAVYSHOTGUN', 'DBSHOTGUN', 'AUTOSHOTGUN', 'COMBATSHOTGUN')
RIFLES = ('ASSAULTRIFLE', 'CARBINERIFLE', 'ADVANCEDRIFLE', 'SPECIALCARBINE', 'BULLPUPRIFLE', 'COMPACTRIFLE', 'MILITARYRIFLE', 'HEAVYRIFLE', 'TACTICALRIFLE', 'BATTLERIFLE')
MG = ('MG', 'COMBATMG', 'GUSENBERG')
SNIPERS = ('SNIPERRIFLE', 'HEAVYSNIPER', 'MARKSMANRIFLE', 'PRECISIONRIFLE')
HEAVY = ('RPG', 'GRENADELAUNCHER', 'MINIGUN', 'FIREWORK', 'RAILGUN', 'HOMINGLAUNCHER', 'COMPACTLAUNCHER', 'RAYMINIGUN', 'EMPLAUNCHER')
THROWN = ('GRENADE', 'BZGAS', 'SMOKEGRENADE', 'MOLOTOV', 'STICKYBOMB', 'PROXMINE', 'SNOWBALL', 'PIPEBOMB', 'BALL', 'FLARE')
TOOLS = ('PETROLCAN', 'FIREEXTINGUISHER', 'HAZARDCAN', 'FERTILIZERCAN', 'ACIDPACKAGE', 'METALDETECTOR', 'DIGISCANNER', 'PARACHUTE')

def categorie(code):
    base = code[len('WEAPON_'):]
    core = base[:-4] if base.endswith('_MK2') else base
    mk2 = ' (Mk II)' if base.endswith('_MK2') else ''
    for names, label in ((MELEE, 'Arme blanche'), (PISTOLS, 'Pistolet'), (SMG, 'Pistolet-mitrailleur'), (SHOTGUNS, 'Fusil à pompe'),
                         (RIFLES, "Fusil d'assaut"), (MG, 'Mitrailleuse'), (SNIPERS, 'Fusil de précision'), (HEAVY, 'Arme lourde'),
                         (THROWN, 'Projectile'), (TOOLS, 'Outil')):
        if core in names: return label + mk2
    return 'Autre'

def parse_lua_table(text, name):
    """Entrées « ['clé'] = { label = …, weight = …, ammoname = … } » d'une table Lua nommée."""
    start = text.index(name + ' = {')
    depth, i = 0, start + len(name) + 3
    while i < len(text):
        if text[i] == '{': depth += 1
        elif text[i] == '}':
            depth -= 1
            if depth == 0: break
        i += 1
    block = text[start:i]
    out = {}
    for m in re.finditer(r"\['([A-Za-z0-9_\-]+)'\]\s*=\s*\{(.*?)\n\t\t\}", block, flags=re.S):
        body = m.group(2)
        g = lambda k: (re.search(k + r"\s*=\s*'([^']*)'", body) or re.search(k + r"\s*=\s*([0-9.]+)", body))
        lab, w, am = g('label'), g('weight'), g('ammoname')
        out[m.group(1)] = {'label': lab.group(1) if lab else '', 'weight': float(w.group(1)) if w else None, 'ammo': am.group(1) if am else ''}
    return out

def armes(ox_weapons_path):
    fr = json.load(io.open('server/locales-fr/weapons.json', encoding='utf-8'))
    ox = {'Weapons': {}, 'Components': {}, 'Ammo': {}}
    if ox_weapons_path:
        txt = io.open(ox_weapons_path, encoding='utf-8').read()
        for k in ox: ox[k] = parse_lua_table(txt, k)
    # Ammu-Nation (vraie boutique RoadLine) et marché noir
    legal = {}
    shops = io.open('server/overrides/ox_inventory/data/shops.lua', encoding='utf-8').read()
    amm = shops[shops.index('Ammunation = {'):]
    amm = amm[:amm.index('locations = {')]
    for m in re.finditer(r"\{\s*name = '([^']+)',\s*price = (\d+)(.*?)\}", amm):
        lic = re.search(r"license = '(\w+)'", m.group(3))
        legal[m.group(1)] = (int(m.group(2)), {'weapon': 'permis de port d\'arme', 'hunting': 'permis de chasse'}.get(lic.group(1) if lic else '', 'sans permis'))
    black = {}
    bm = io.open('server/resources/[gtasoon]/gs_blackmarket/shared/config.lua', encoding='utf-8').read()
    for m in re.finditer(r"\{\s*item = '([^']+)',\s*label = '([^']*)',(.*?)\}", bm):
        price = re.search(r"price = (\d+)", m.group(3)); stock = re.search(r"stock = (\d+)", m.group(3))
        black[m.group(1)] = (int(price.group(1)) if price else None, int(stock.group(1)) if stock else None, 'gangOnly = true' in m.group(3))
    def vente(code):
        parts = []
        if code in legal: parts.append(f'Ammu-Nation {legal[code][0]:,} $ ({legal[code][1]})'.replace(',', ' '))
        if code in black:
            p, s, g = black[code]
            parts.append(f'Marché noir {p:,} $ sale'.replace(',', ' ') + (f', stock {s}/jour' if s else '') + (' · gang seulement' if g else ''))
        return ' · '.join(parts) if parts else 'Non vendu (saisie, police, staff)'
    def prix(code):
        if code in legal: return legal[code][0]
        if code in black and black[code][0]: return black[code][0]
        return None
    rows_w, rows_a, rows_c = [], [], []
    for code, label in sorted(fr.items(), key=lambda kv: (0 if kv[0].startswith('WEAPON_') else 1 if kv[0].startswith('ammo') else 2, kv[1])):
        o = ox['Weapons'].get(code) or ox['Ammo'].get(code) or ox['Components'].get(code) or {}
        w = o.get('weight'); kg = round(w / 1000, 2) if w else None
        if code.startswith('WEAPON_'):
            am = o.get('ammo', '')
            rows_w.append((label, code, categorie(code), fr.get(am, am) if am else '—', kg, vente(code), prix(code)))
        elif code.startswith('ammo'):
            users = sorted(fr.get(k, k) for k, v in ox['Weapons'].items() if v.get('ammo') == code)
            rows_a.append((label, code, ', '.join(users) if users else '—', kg, vente(code), prix(code)))
        else:
            rows_c.append((label, code, kg, vente(code), prix(code)))
    rows_w.sort(key=lambda r: (r[2], r[0]))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Armes'
    sheet_table(ws, ['Arme', 'Code', 'Catégorie', 'Munition', 'Poids (kg)', 'Où l\'acheter', 'Prix'], rows_w, [30, 30, 22, 26, 11, 58, 11], money_cols=(7,),
                note='Noms : server/locales-fr/weapons.json · prix : Ammu-Nation (server/overrides/ox_inventory/data/shops.lua, permis exigé) et marché noir '
                     '(gs_blackmarket, argent sale, nuit, stock par jour, +40 % en liquide propre) · poids : ox_inventory.')
    ws2 = wb.create_sheet('Munitions')
    sheet_table(ws2, ['Munition', 'Code', 'Armes qui l\'utilisent', 'Poids (kg)', 'Où l\'acheter', 'Prix'], rows_a, [26, 16, 60, 11, 50, 11], money_cols=(6,),
                note='Ammu-Nation : 120 munitions par jour et par personnage, permis de port d\'arme exigé (gs_blackmarket Config.Legal).')
    ws3 = wb.create_sheet('Accessoires')
    sheet_table(ws3, ['Accessoire', 'Code', 'Poids (kg)', 'Où l\'acheter', 'Prix'], rows_c, [34, 30, 11, 50, 11], money_cols=(5,))
    n = len(rows_w)
    lines = [('Armes', f'=COUNTA(Armes!B2:B{n + 1})', '0'), ('Munitions', f'=COUNTA(Munitions!B2:B{len(rows_a) + 1})', '0'),
             ('Accessoires', f'=COUNTA(Accessoires!B2:B{len(rows_c) + 1})', '0'),
             ('Armes vendues à l\'Ammu-Nation', f'=COUNTIF(Armes!F2:F{n + 1},"Ammu-Nation*")', '0'),
             ('Armes vendues au marché noir', f'=COUNTIF(Armes!F2:F{n + 1},"*Marché noir*")', '0'),
             ('Armes non vendues (saisie, police, staff)', f'=COUNTIF(Armes!F2:F{n + 1},"Non vendu*")', '0'), ('', None, None)]
    for c in sorted(set(r[2] for r in rows_w)):
        lines.append((c, f'=COUNTIF(Armes!C2:C{n + 1},"{c}")', '0'))
    summary_sheet(wb, 'Résumé', 'RoadLine RP · Armes, munitions et accessoires (résumé)', lines)
    wb.calculation.fullCalcOnLoad = True
    wb.save('docs/ROADLINE_Armes.xlsx')
    return len(rows_w), len(rows_a), len(rows_c)

if __name__ == '__main__':
    nv = vehicules()
    na = armes(sys.argv[1] if len(sys.argv) > 1 else None)
    print(f'docs/ROADLINE_Vehicules.xlsx ({nv} véhicules), docs/ROADLINE_Armes.xlsx ({na[0]} armes, {na[1]} munitions, {na[2]} accessoires)')
